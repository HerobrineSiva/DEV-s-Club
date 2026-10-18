"""Build the final Track 1 Word report and its original visual assets."""

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "Machine_Failure_Prediction_Hackathon_Report.docx"
ASSETS = ROOT / "report_assets"

NAVY = "17365D"
BLUE = "2F75B5"
PALE = "EAF2F8"
LIGHT = "F5F7FA"
GRAY = "D9E2F3"
RED = "C00000"
GREEN = "2E7D32"


def font(size, bold=False):
    try:
        return ImageFont.truetype("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf", size)
    except OSError:
        return ImageFont.load_default()


def save_visuals():
    ASSETS.mkdir(exist_ok=True)
    # Cover illustration: original abstract CNC/sensor visualization.
    im = Image.new("RGB", (1600, 760), "#0B1F33")
    d = ImageDraw.Draw(im)
    for x in range(0, 1600, 80): d.line((x, 0, x, 760), fill="#123454", width=1)
    for y in range(0, 760, 80): d.line((0, y, 1600, y), fill="#123454", width=1)
    d.rounded_rectangle((150, 200, 850, 580), radius=35, fill="#EAF2F8", outline="#5DADE2", width=8)
    d.rectangle((330, 100, 640, 280), fill="#C8D6E5", outline="#5DADE2", width=8)
    d.rectangle((420, 280, 540, 450), fill="#85929E")
    d.ellipse((395, 430, 565, 600), fill="#2F75B5", outline="#B8E0F5", width=8)
    d.line((480, 600, 480, 680), fill="#B8E0F5", width=20)
    d.polygon([(480, 680), (440, 735), (520, 735)], fill="#B8E0F5")
    for cx, cy in [(1000,180),(1250,260),(1110,510),(1420,560)]:
        d.ellipse((cx-30, cy-30, cx+30, cy+30), fill="#F4D03F")
        d.line((850, 390, cx, cy), fill="#5DADE2", width=6)
    d.text((920, 330), "Telemetry", font=font(46, True), fill="white")
    d.text((920, 390), "Risk scoring", font=font(46, True), fill="white")
    im.save(ASSETS / "cover_visual.png")

    # Class distribution.
    im = Image.new("RGB", (1300, 650), "white"); d = ImageDraw.Draw(im)
    d.text((55, 35), "Class distribution in the training data", font=font(42, True), fill="#17365D")
    labels = [("Healthy", 7728, "#4F81BD"), ("Failure", 271, "#C0504D")]
    for i,(label,value,color) in enumerate(labels):
        y=170+i*190; width=int(value/7728*700)
        d.rounded_rectangle((250,y,250+width,y+95), 15, fill=color)
        d.text((55,y+20), label, font=font(34,True), fill="#17202A")
        d.text((260+width,y+20), f"{value:,}  ({value/7999:.2%})", font=font(30), fill="#17202A")
    d.text((55,570), "Rare failures make PR-AUC and F1 more informative than accuracy.", font=font(28), fill="#566573")
    im.save(ASSETS / "class_distribution.png")

    # AP comparison.
    im = Image.new("RGB", (1300, 760), "white"); d = ImageDraw.Draw(im)
    d.text((55, 30), "Development out-of-fold Average Precision", font=font(42, True), fill="#17365D")
    models=[("Logistic Regression",.4139,"#95A5A6"),("Random Forest",.7199,"#5DADE2"),("Extra Trees",.6413,"#7FB3D5"),("HistGradientBoosting",.7853,"#2E7D32")]
    x0=360; scale=850
    for i,(name,val,color) in enumerate(models):
        y=145+i*130
        d.text((40,y+22),name,font=font(27, True),fill="#17202A")
        d.rounded_rectangle((x0,y,x0+int(val*scale),y+70),12,fill=color)
        d.text((x0+int(val*scale)+20,y+18),f"{val:.4f}",font=font(28,True),fill="#17202A")
    d.line((x0,690,x0+scale,690), fill="#566573", width=3)
    for tick in [0,.2,.4,.6,.8,1.0]: d.text((x0+int(tick*scale)-12,700),str(tick),font=font(22),fill="#566573")
    im.save(ASSETS / "model_comparison.png")

    # Confusion matrix.
    im=Image.new("RGB",(900,760),"white"); d=ImageDraw.Draw(im)
    d.text((80,35),"Final holdout confusion matrix",font=font(42,True),fill="#17365D")
    d.text((330,110),"Predicted healthy",font=font(23,True),fill="#17202A")
    d.text((595,110),"Predicted failure",font=font(23,True),fill="#17202A")
    d.text((45,265),"Actual healthy",font=font(24,True),fill="#17202A")
    d.text((45,455),"Actual failure",font=font(24,True),fill="#17202A")
    cells=[(280,160,1527,"#D5E8F7"),(535,160,19,"#F9E0DF"),(280,350,13,"#F9E0DF"),(535,350,41,"#D9EAD3")]
    for x,y,val,color in cells:
        d.rectangle((x,y,x+220,y+160),fill=color,outline="#AAB7B8",width=3)
        d.text((x+72,y+56),str(val),font=font(44,True),fill="#17202A")
    d.text((80,610),"41 of 54 failures detected; 19 healthy machines flagged.",font=font(28),fill="#566573")
    im.save(ASSETS / "confusion_matrix.png")

    # Workflow visual.
    im=Image.new("RGB",(1500,430),"white"); d=ImageDraw.Draw(im)
    boxes=[("Train data", "6 telemetry features\n+failure label"),("Stratified split", "Development 80%\nHoldout 20%"),("Five-fold OOF CV", "Compare four models\nrank by PR-AUC"),("Threshold selection", "F1 maximized on\ndevelopment only"),("Final artifact", "Probability +\nclass prediction")]
    for i,(title,body) in enumerate(boxes):
        x=35+i*295; d.rounded_rectangle((x,95,x+235,310),20,fill="#EAF2F8",outline="#2F75B5",width=4)
        d.text((x+25,125),title,font=font(25,True),fill="#17365D")
        d.multiline_text((x+25,180),body,font=font(22),fill="#17202A",spacing=8)
        if i<4: d.polygon([(x+250,190),(x+280,205),(x+250,220)],fill="#2F75B5")
    d.text((35,28),"Reproducible machine failure prediction workflow",font=font(38,True),fill="#17365D")
    im.save(ASSETS / "workflow.png")


def shade(cell, color):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),color); tcPr.append(shd)


def border_table(table):
    tblPr=table._tbl.tblPr; borders=OxmlElement("w:tblBorders")
    for edge in ("top","left","bottom","right","insideH","insideV"):
        el=OxmlElement(f"w:{edge}"); el.set(qn("w:val"),"single"); el.set(qn("w:sz"),"6"); el.set(qn("w:color"),"D9D9D9"); borders.append(el)
    tblPr.append(borders)


def add_heading(doc, text, level=1):
    p=doc.add_paragraph(style=f"Heading {level}"); p.paragraph_format.space_before=Pt(16); p.paragraph_format.space_after=Pt(7)
    r=p.add_run(text); r.font.color.rgb=RGBColor(0,0,0); return p


def add_text(doc, text, bold_start=None):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(7); p.paragraph_format.line_spacing=1.16
    if bold_start and text.startswith(bold_start):
        p.add_run(bold_start).bold=True; p.add_run(text[len(bold_start):])
    else: p.add_run(text)
    return p


def add_caption(doc, text):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(11)
    r=p.add_run(text); r.italic=True; r.font.size=Pt(9); r.font.color.rgb=RGBColor(89,89,89)


def add_hyperlink(paragraph, text, url):
    part=paragraph.part; rid=part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink=OxmlElement("w:hyperlink"); hyperlink.set(qn("r:id"),rid)
    new_run=OxmlElement("w:r"); rPr=OxmlElement("w:rPr"); color=OxmlElement("w:color"); color.set(qn("w:val"),"0563C1"); rPr.append(color); u=OxmlElement("w:u"); u.set(qn("w:val"),"single"); rPr.append(u); new_run.append(rPr)
    t=OxmlElement("w:t"); t.text=text; new_run.append(t); hyperlink.append(new_run); paragraph._p.append(hyperlink)


def build():
    save_visuals(); doc=Document(); section=doc.sections[0]; section.top_margin=Inches(.7); section.bottom_margin=Inches(.7); section.left_margin=Inches(.75); section.right_margin=Inches(.75)
    styles=doc.styles; styles["Normal"].font.name="Aptos"; styles["Normal"].font.size=Pt(10.5)
    for style in ["Title","Heading 1","Heading 2"]: styles[style].font.name="Aptos"; styles[style].font.color.rgb=RGBColor(0,0,0)
    # Header/footer.
    footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; footer.add_run("MODEL ARENA Track 1 | Machine Failure Prediction").font.size=Pt(8)
    # Cover.
    p=doc.add_paragraph(); p.style="Title"; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run("Machine Failure Prediction for Predictive Maintenance")
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run("MODEL ARENA Hackathon | Track 1 Beginner"); r.bold=True; r.font.size=Pt(16); r.font.color.rgb=RGBColor(23,54,93)
    doc.add_paragraph("")
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(ASSETS / "cover_visual.png"),width=Inches(6.7))
    add_text(doc,"Executive summary: This report presents a reproducible machine-learning system that predicts failure risk from six manufacturing telemetry features. HistGradientBoosting was selected using development-only out-of-fold Average Precision, then evaluated on a fixed local holdout. The resulting system reaches 0.8086 Average Precision and 0.7193 F1, and exports a validated inference artifact for organizer-supplied hidden test data.")
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run("Prepared for hackathon evaluation | October 2026").italic=True
    doc.add_page_break()

    add_heading(doc,"1 Problem Statement and Business Context")
    add_text(doc,"High-precision milling and lathe machines operate continuously. A sudden failure can stop a production line, create downtime, and increase emergency repair costs. Replacing tools on fixed schedules can also waste useful components. The Track 1 task is to predict whether an equipment observation signals an imminent machine failure while keeping false alarms manageable.")
    add_text(doc,"The system is a binary classifier. It receives the product type, temperature conditions, spindle rotation speed, torque, and tool wear. It returns a continuous failure probability and a thresholded failure alert. This design supports both ranking machines by risk and taking a clear inspection decision.")
    add_heading(doc,"2 Dataset and Exploratory Findings")
    add_text(doc,"The supplied training file contains 7,999 labelled observations. It has six input features and the target Machine failure, where 0 means healthy and 1 means failure. Basic inspection found no missing values and no duplicate records.")
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(ASSETS / "class_distribution.png"),width=Inches(6.6)); add_caption(doc,"Figure 1. Training-class distribution. Failures occur in only 3.39% of observations.")
    add_text(doc,"Because failures are rare, accuracy would be misleading. Predicting every machine as healthy would score 96.61% accuracy while detecting no failures. The project therefore uses Average Precision, also called PR-AUC in the hackathon criteria, as the primary ranking measure and uses F1, precision, recall, and the confusion matrix to assess alert decisions.")
    add_heading(doc,"Input Features",2)
    table=doc.add_table(rows=1,cols=3); table.style="Table Grid"; border_table(table)
    for cell,text in zip(table.rows[0].cells,["Feature","Meaning","Validation"]): shade(cell,NAVY); cell.text=text; cell.vertical_alignment=WD_ALIGN_VERTICAL.CENTER
    for cell in table.rows[0].cells:
        for run in cell.paragraphs[0].runs: run.font.color.rgb=RGBColor(255,255,255); run.bold=True
    rows=[("Type","Component quality variant: L, M, or H","Must be L, M, or H"),("Air and process temperature","Ambient and enclosure temperature in Kelvin","Finite numeric values"),("Rotational speed","Spindle speed in rpm","Positive numeric value"),("Torque","Resistance torque in Nm","Finite numeric value"),("Tool wear","Cumulative cutting-tool use in minutes","Non-negative numeric value")]
    for idx,row in enumerate(rows):
        cells=table.add_row().cells
        for c,text in zip(cells,row): c.text=text; c.vertical_alignment=WD_ALIGN_VERTICAL.CENTER; shade(c,LIGHT if idx%2==0 else "FFFFFF")
    doc.add_page_break()

    add_heading(doc,"3 Modelling Methodology")
    add_text(doc,"All preprocessing is implemented inside sklearn Pipeline and ColumnTransformer objects. This is important because it stops validation information from leaking into training. Numeric fields use median imputation; Logistic Regression additionally uses standard scaling; the Type category is imputed and one-hot encoded. Tree models do not require scaling, while HistGradientBoosting receives dense encoded input because of its estimator requirements.")
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(ASSETS / "workflow.png"),width=Inches(6.8)); add_caption(doc,"Figure 2. The reproducible experiment and deployment workflow.")
    add_heading(doc,"Validation Design",2)
    add_text(doc,"A fixed 80/20 stratified split with random seed 42 reserves 1,600 observations, including 54 failures, as a local holdout. The remaining 6,399 development observations are used for five-fold StratifiedKFold cross-validation. Every candidate model is evaluated against the same development/holdout boundary.")
    add_text(doc,"Out-of-fold probabilities are created for each development row. The best model is chosen by development OOF Average Precision. For the selected model only, thresholds from 0.05 through 0.95 are searched and the one with the best OOF F1 is retained. The holdout labels do not participate in choosing the model or threshold.")
    add_heading(doc,"Models Compared",2)
    models=doc.add_table(rows=1,cols=3); models.style="Table Grid"; border_table(models)
    for cell,text in zip(models.rows[0].cells,["Model","Reason for inclusion","Class imbalance handling"]): shade(cell,NAVY); cell.text=text
    for cell in models.rows[0].cells:
        for run in cell.paragraphs[0].runs: run.font.color.rgb=RGBColor(255,255,255); run.bold=True
    for idx,row in enumerate([("Logistic Regression","Interpretable linear baseline","balanced class weights"),("Random Forest","Nonlinear ensemble baseline","balanced subsample weights"),("Extra Trees","Randomized tree ensemble","balanced class weights"),("HistGradientBoosting","Efficient gradient-boosted trees","selection based on PR-AUC")]):
        cells=models.add_row().cells
        for c,text in zip(cells,row): c.text=text; shade(c,LIGHT if idx%2==0 else "FFFFFF")
    doc.add_page_break()

    add_heading(doc,"4 Model Comparison and Threshold Selection")
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(ASSETS / "model_comparison.png"),width=Inches(6.7)); add_caption(doc,"Figure 3. Development out-of-fold Average Precision across candidate models.")
    result=doc.add_table(rows=1,cols=4); result.style="Table Grid"; border_table(result)
    headers=["Model","OOF AP","F1 threshold","OOF F1"]
    for cell,text in zip(result.rows[0].cells,headers): shade(cell,NAVY); cell.text=text
    for cell in result.rows[0].cells:
        for run in cell.paragraphs[0].runs: run.font.color.rgb=RGBColor(255,255,255); run.bold=True
    vals=[("Logistic Regression","0.4139","0.88","0.4350"),("Random Forest","0.7199","0.35","0.6809"),("Extra Trees","0.6413","0.54","0.5982"),("HistGradientBoosting","0.7853","0.26","0.7378")]
    for idx,row in enumerate(vals):
        cells=result.add_row().cells
        for c,text in zip(cells,row): c.text=text; shade(c,"D9EAD3" if idx==3 else (LIGHT if idx%2==0 else "FFFFFF")); c.vertical_alignment=WD_ALIGN_VERTICAL.CENTER
    add_text(doc,"HistGradientBoosting achieved the strongest development OOF Average Precision (0.7853), so it was selected. Its F1-optimal threshold was 0.26. This is below the usual 0.50 classification threshold because a failure-alert setting benefits from identifying more rare failures, while the F1 criterion balances that additional recall against false alerts.")
    add_heading(doc,"Why PR-AUC and F1",2)
    add_text(doc,"PR-AUC evaluates how well the probability ranking distinguishes rare positive cases across thresholds. F1 combines precision and recall at one chosen threshold. These measures match the hackathon rules and are more informative than accuracy for a failure rate of 3.39%.")
    p=doc.add_paragraph(); p.add_run("Web reference: ").bold=True; add_hyperlink(p,"NIST explains that a precision-recall curve is formed by varying the confidence threshold and relates precision to positive predictive value and recall to sensitivity.","https://nvlpubs.nist.gov/nistpubs/ir/2014/nist.ir.7972.pdf")
    doc.add_page_break()

    add_heading(doc,"5 Final Holdout Results")
    add_text(doc,"After model and threshold selection, the HistGradientBoosting pipeline was fitted only on the development partition and evaluated once on the fixed 1,600-row holdout. This is the clearest local estimate of expected generalization for this split.")
    metrics=doc.add_table(rows=1,cols=2); metrics.style="Table Grid"; border_table(metrics)
    for cell,text in zip(metrics.rows[0].cells,["Metric","Holdout result"]): shade(cell,NAVY); cell.text=text
    for cell in metrics.rows[0].cells:
        for run in cell.paragraphs[0].runs: run.font.color.rgb=RGBColor(255,255,255); run.bold=True
    for idx,row in enumerate([("Holdout observations","1,600"),("Actual failures","54"),("Average Precision / PR-AUC","0.8086"),("F1-score","0.7193"),("Precision","0.6833"),("Recall","0.7593"),("Decision threshold","0.26")]):
        cells=metrics.add_row().cells
        for c,text in zip(cells,row): c.text=text; shade(c,LIGHT if idx%2==0 else "FFFFFF")
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(ASSETS / "confusion_matrix.png"),width=Inches(5.7)); add_caption(doc,"Figure 4. Local holdout confusion matrix at the selected threshold of 0.26.")
    add_text(doc,"The model detected 41 of the 54 failure cases and missed 13. It also flagged 19 healthy machines. For an industrial team, these outputs are actionable: the continuous probability can rank machines by risk, while the alert class can trigger an inspection workflow.")
    doc.add_page_break()

    add_heading(doc,"6 Implementation and Submission Workflow")
    add_text(doc,"The codebase is organized so that training and inference use the same feature schema and preprocessing. The exported joblib artifact stores the fitted model pipeline, feature order, model name, selected threshold, positive class, random seed, and training metadata.")
    add_heading(doc,"Key Commands",2)
    code=doc.add_table(rows=1,cols=1); border_table(code); shade(code.rows[0].cells[0],"F2F2F2")
    code.rows[0].cells[0].text=("# Install dependencies\npython -m pip install -r requirements.txt\n\n# Reproduce selection and local evaluation\npython src\\run_experiment.py\n\n# Fit the deployment artifact after evaluation\npython src\\train_final.py\n\n# Score an organizer-supplied hidden test file\npython src\\make_submission.py path\\to\\test.csv --output submission.csv\n\n# Run automated checks\npython -m pytest -q")
    for run in code.rows[0].cells[0].paragraphs[0].runs: run.font.name="Consolas"; run.font.size=Pt(9)
    add_heading(doc,"Example Input and Output",2)
    ex=doc.add_table(rows=1,cols=2); ex.style="Table Grid"; border_table(ex)
    for cell,text in zip(ex.rows[0].cells,["Example input","Model output"]): shade(cell,NAVY); cell.text=text
    for cell in ex.rows[0].cells:
        for run in cell.paragraphs[0].runs: run.font.color.rgb=RGBColor(255,255,255); run.bold=True
    example_row = ex.add_row().cells
    example_row[0].text="Type: M\nAir temperature: 300.1 K\nProcess temperature: 310.2 K\nSpeed: 1500 rpm\nTorque: 40.5 Nm\nTool wear: 120 min"
    example_row[1].text="failure_probability: 0.00096\npredicted_failure: 0\n\nThe actual result depends on the supplied telemetry."
    add_text(doc,"For a hidden organizer test set, the input must contain the same six feature columns and no target column. The submission script produces both Machine failure probability, for PR-AUC scoring, and Machine failure, for the thresholded decision. The organizer's exact column naming rule should be checked before upload.")
    add_heading(doc,"7 Limitations and Next Steps")
    add_text(doc,"The local result is strong for this specific random stratified split, but the dataset contains no timestamp, machine identifier, site identifier, or maintenance-cost information. The evaluation therefore cannot prove temporal robustness, performance on entirely new machines, or cost-optimal behavior. The selected 0.26 threshold maximizes F1 locally; a production maintenance policy should adjust it using the relative cost of missed failures, downtime, inspections, and false alarms.")
    add_text(doc,"Recommended next steps are to collect time-ordered and machine-identified telemetry, validate with temporal or group-aware splits, analyze probability calibration, and select the alert threshold jointly with maintenance stakeholders.")
    add_heading(doc,"References")
    p=doc.add_paragraph(style=None); p.add_run("1. "); add_hyperlink(p,"NIST IR 7972, Evaluation of Pattern Recognition Algorithms for Machine Learning, discussion of precision-recall curves and threshold variation.","https://nvlpubs.nist.gov/nistpubs/ir/2014/nist.ir.7972.pdf")
    p=doc.add_paragraph(); p.add_run("2. "); add_hyperlink(p,"scikit-learn model evaluation documentation, precision-recall tools and metrics.","https://scikit-learn.org/stable/modules/model_evaluation.html")
    add_text(doc,"All charts and diagrams in this report were created from the project dataset summary and measured local experiment results. The web references provide metric context; no external images were reproduced.")
    doc.core_properties.title="Machine Failure Prediction for Predictive Maintenance"; doc.core_properties.author="DEVS Club"; doc.save(OUT)
    print(OUT)


if __name__ == "__main__": build()
